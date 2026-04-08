"""
이배 《En attendant: 기다리며》 도슨트 PPT 빌드 스크립트
images/ 폴더의 이미지를 사용합니다.
출력: leebae_docent_final.pptx
"""

import os, sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from lxml import etree
except ImportError:
    print("필요 패키지 설치: pip install python-pptx lxml")
    sys.exit(1)

SLIDE_W  = Emu(9144000)
SLIDE_H  = Emu(5143500)
C_BG     = RGBColor(0x11, 0x11, 0x11)
C_GOLD   = RGBColor(0xC8, 0xA8, 0x4B)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_GRAY_L = RGBColor(0xBB, 0xBB, 0xBB)
C_GRAY_M = RGBColor(0x88, 0x88, 0x88)
C_GRAY_D = RGBColor(0x4A, 0x4A, 0x4A)
C_CARD   = RGBColor(0x22, 0x22, 0x22)
C_CARD_B = RGBColor(0x3A, 0x3A, 0x3A)

def load_images():
    imgs = {}
    img_dir = Path("images")
    for key in ["img_1","img_2","img_3","img_4","img_5","img_6","img_7"]:
        p = img_dir / f"{key}.jpg"
        imgs[key] = p if p.exists() else None
        status = "OK" if p.exists() else "없음"
        print(f"  {key}: {status}")
    return imgs

def _hex(c): return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"

def set_bg_color(slide, color):
    cSld = slide._element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg"))
    if bg is None:
        bg = etree.SubElement(cSld, qn("p:bg"))
        cSld.insert(0, bg)
    bgPr = bg.find(qn("p:bgPr"))
    if bgPr is None:
        bgPr = etree.SubElement(bg, qn("p:bgPr"))
    for c in list(bgPr): bgPr.remove(c)
    clr = etree.SubElement(etree.SubElement(bgPr, qn("a:solidFill")), qn("a:srgbClr"))
    clr.set("val", _hex(color))

def set_bg_image(slide, img_path):
    if img_path is None:
        return set_bg_color(slide, C_BG)
    _, rId = slide.part.get_or_add_image_part(str(img_path))
    cSld = slide._element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg"))
    if bg is None:
        bg = etree.SubElement(cSld, qn("p:bg")); cSld.insert(0, bg)
    bgPr = bg.find(qn("p:bgPr"))
    if bgPr is None: bgPr = etree.SubElement(bg, qn("p:bgPr"))
    for c in list(bgPr): bgPr.remove(c)
    bf = etree.SubElement(bgPr, qn("a:blipFill"))
    bf.set("dpi","0"); bf.set("rotWithShape","1")
    blip = etree.SubElement(bf, qn("a:blip"))
    blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rId)
    etree.SubElement(etree.SubElement(bf, qn("a:stretch")), qn("a:fillRect"))

def add_img(slide, p, x, y, w, h):
    if p: slide.shapes.add_picture(str(p), Emu(x), Emu(y), Emu(w), Emu(h))

def add_rect(slide, x, y, w, h, fill=None, fa=None, line=None, lw=12700):
    s = slide.shapes.add_shape(1, Emu(x), Emu(y), Emu(w), Emu(h))
    sp = s.element.find(qn("p:spPr"))
    for t in [qn("a:solidFill"), qn("a:noFill")]:
        for e in sp.findall(t): sp.remove(e)
    if fill:
        clr = etree.SubElement(etree.SubElement(sp, qn("a:solidFill")), qn("a:srgbClr"))
        clr.set("val", _hex(fill))
        if fa: etree.SubElement(clr, qn("a:alpha")).set("val", str(fa))
    else: etree.SubElement(sp, qn("a:noFill"))
    ln = sp.find(qn("a:ln"))
    if ln is None: ln = etree.SubElement(sp, qn("a:ln"))
    for e in list(ln): ln.remove(e)
    if line:
        ln.set("w", str(lw))
        clr2 = etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"))
        clr2.set("val", _hex(line))
    else: etree.SubElement(ln, qn("a:noFill"))

def _bpr(tf, anc="t"):
    b = tf._txBody.find(qn("a:bodyPr"))
    if b is not None:
        b.set("anchor", anc)
        for a in ["lIns","tIns","rIns","bIns"]: b.set(a, "0")

def add_text(slide, x, y, w, h, text, sz, color,
             bold=False, italic=False, font="Calibri",
             align=PP_ALIGN.LEFT, anchor="t", spc=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True; _bpr(tf, anchor)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    f = r.font; f.size=Pt(sz); f.bold=bold; f.italic=italic
    f.name=font; f.color.rgb=color
    if spc:
        rp = r._r.find(qn("a:rPr"))
        if rp is None: rp = etree.SubElement(r._r, qn("a:rPr"))
        rp.set("spc", str(spc))

def add_lines(slide, x, y, w, h, lines, sz, color,
              bold=False, italic=False, font="Calibri",
              align=PP_ALIGN.LEFT, anchor="t"):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True; _bpr(tf, anchor)
    for i, line in enumerate(lines):
        if isinstance(line, str): txt, lb, lc = line, bold, color
        else: txt, lb, lc = line
        p = tf.paragraphs[i] if i==0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = txt
        f = r.font; f.size=Pt(sz); f.bold=lb; f.italic=italic
        f.name=font; f.color.rgb=lc

def add_card(slide, x, y, w, h, icon, title, sub, body):
    P = 182880
    add_rect(slide, x, y, w, h, fill=C_CARD, line=C_CARD_B)
    add_rect(slide, x, y, w, 54864, fill=C_GOLD)
    cy = y+120960
    add_text(slide, x+P, cy, w-P*2, 320040, f"{icon}  {title}", 13, C_WHITE, bold=True)
    cy += 342000
    if sub:
        add_text(slide, x+P, cy, w-P*2, 260000, sub, 10, C_GOLD, italic=True)
        cy += 285000
    add_text(slide, x+P, cy, w-P*2, h-(cy-y)-P, body, 11, C_GRAY_L)

# ──────────────────────────────────────────────
# 슬라이드 1 — 표지
# ──────────────────────────────────────────────
def s1(slide, imgs):
    set_bg_color(slide, C_BG)
    add_img(slide, imgs["img_1"], 4500000, 0, 4644000, 5143500)
    add_rect(slide, 4500000, 0, 4644000, 5143500, fill=RGBColor(0x11,0x11,0x11), fa=60000)
    add_rect(slide, 411480, 731520, 54864, 3657600, fill=C_GOLD)
    add_text(slide, 640080, 731520, 8000000, 1005840, "이배 李培",
             52, C_WHITE, bold=True, font="Georgia", anchor="ctr")
    add_text(slide, 640080, 1554480, 8000000, 457200, "L E E   B A E",
             18, C_GOLD, font="Georgia", spc=800, anchor="ctr")
    add_text(slide, 640080, 2286000, 7700000, 640080, "En attendant: 기다리며",
             26, C_GRAY_L, italic=True, anchor="ctr")
    add_lines(slide, 640080, 3108960, 7300000, 731520,
              ["뮤지엄 산 (Museum SAN), 강원도 원주", "2026. 4. 7 — 12. 6"],
              14, C_GRAY_M)
    add_text(slide, 0, 4572000, 9144000, 365760,
             "뮤지엄 산 최초 국내 작가 기획전  ·  역대 최대 규모  ·  39점 전시",
             11, C_GRAY_D, align=PP_ALIGN.CENTER, anchor="ctr")
    for bx, by, bh, ba in [
        (7772400,274320,4114800,5000),(7452360,548640,3749040,9000),
        (7132320,822960,3383280,13000),(6812280,1097280,3017520,17000),
        (6492240,1371600,2651760,21000),(6172200,1645920,2286000,25000)]:
        add_rect(slide, bx, by, 1097280, bh, fill=C_GRAY_D, fa=ba, line=C_GRAY_D, lw=12700)

# ──────────────────────────────────────────────
# 슬라이드 2 — 작가 소개
# ──────────────────────────────────────────────
def s2(slide, imgs):
    set_bg_color(slide, C_BG)
    add_img(slide, imgs["img_7"], 6100000, 300000, 3044000, 4600000)
    add_rect(slide, 0, 0, 9144000, 502920, fill=RGBColor(0x1A,0x1A,0x1A))
    add_text(slide, 457200, 0, 8229600, 502920, "A R T I S T",
             11, C_GRAY_M, spc=600, anchor="ctr")
    add_text(slide, 457200, 685800, 5500000, 822960,
             "숯의 작가 이배", 36, C_WHITE, bold=True)
    add_rect(slide, 365760, 1645920, 3657600, 3108960,
             fill=RGBColor(0x22,0x22,0x22), line=RGBColor(0x3A,0x3A,0x3A))
    for i, (label, val) in enumerate([
        ("출생","1956년, 경상북도 청도"),("학력","홍익대학교 서양화"),
        ("도불","1989년 프랑스 파리 정착"),("거점","파리 · 청도 · 고양"),("나이","70세")]):
        ry = 1828800 + i * 530000
        add_text(slide, 548640, ry, 914400, 365760, label, 13, C_GOLD, bold=True)
        add_text(slide, 1554480, ry, 2400000, 365760, val, 13, C_GRAY_L)
    add_text(slide, 4389120, 1645920, 1600000, 365760,
             "농부의 아들에서 세계적 작가로", 13, C_WHITE, bold=True)
    add_lines(slide, 4389120, 2100000, 1600000, 2654880,
              ["경북 청도의 농부 아들로 태어나 1989년 프랑스로 건너간 이배는, "
               "재정적 어려움 속에서 우연히 발견한 '숯 포대' 하나로 자신만의 예술 언어를 찾았습니다.",
               "",
               "이후 30여 년간 숯이라는 단 하나의 매체에 천착하며 파리·뉴욕·베니스·서울 등 "
               "세계 무대에서 인정받는 작가가 되었습니다.",
               "",
               "프랑스 국립 기메 동양박물관 개인전, 뉴욕 록펠러센터 전시, "
               "베니스 빌모트 파운데이션 개인전 등 굵직한 커리어를 이어오고 있습니다."],
              12, C_GRAY_L)

# ──────────────────────────────────────────────
# 슬라이드 3 — 숯의 언어
# ──────────────────────────────────────────────
def s3(slide, imgs):
    set_bg_image(slide, imgs["img_2"])
    add_rect(slide, 0, 0, 9144000, 502920, fill=RGBColor(0x0D,0x0D,0x0D), fa=90000)
    add_text(slide, 457200, 0, 8229600, 502920, "M E D I U M",
             11, C_GRAY_M, spc=600, anchor="ctr")
    add_text(slide, 457200, 685800, 8229600, 822960,
             "숯 — 이배의 언어", 36, C_WHITE, bold=True)
    cw, gap = 2743200, 228600
    sx = (9144000 - (cw*3 + gap*2)) // 2
    for i, (icon, title, sub, body) in enumerate([
        ("🔥","불로부터","Issu du feu",
         "나무는 불 속에서 형태를 잃지만 새로운 물질로 재탄생합니다. 파괴와 재생의 순환."),
        ("⬛","검정의 깊이","검정 속 백 가지 색",
         '"모든 색을 흡수한 검정에는 한 가지 빛깔이 아닌 백 가지의 색이 들어 있다."'),
        ("🌿","정화와 포용","한국적 상징",
         "숯은 한국 전통에서 정화와 치유의 상징입니다. 이배에게 숯은 자신의 뿌리로 돌아가는 길."),
    ]): add_card(slide, sx+i*(cw+gap), 1691640, cw, 3200000, icon, title, sub, body)
    add_text(slide, 0, 4754880, 9144000, 384048,
             '"숯은 나에게 자신의 원천을 일깨워주는 재질이었다." — 이배',
             11, C_GRAY_D, italic=True, align=PP_ALIGN.CENTER, anchor="ctr")

# ──────────────────────────────────────────────
# 슬라이드 4 — 전시 공간
# ──────────────────────────────────────────────
def s4(slide, imgs):
    set_bg_image(slide, imgs["img_5"])
    add_rect(slide, 0, 0, 9144000, 502920, fill=RGBColor(0x0D,0x0D,0x0D), fa=92000)
    add_text(slide, 457200, 0, 8229600, 502920,
             "E X H I B I T I O N   S P A C E S", 11, C_GRAY_M, spc=600, anchor="ctr")
    add_rect(slide, 0, 502920, 9144000, 4640580, fill=RGBColor(0x0A,0x0A,0x0A), fa=75000)
    add_text(slide, 457200, 685800, 8229600, 914400,
             "전시 공간 — 6개의 사유의 장", 40, C_WHITE, bold=True)
    cw, ch, gx, gy = 2743200, 1280160, 228600, 182880
    sx = (9144000 - (cw*3 + gx*2)) // 2
    for i, (title, sub, num, desc) in enumerate([
        ("불로부터","Issu du feu","01","높이 8m · 폭 5m · 무게 7톤\n2023 뉴욕 록펠러센터 화제작의 확장"),
        ("붓질 연작","Brushstroke × 16","02","자연광과 함께 시시각각 변화\n풍경 속을 산책하는 경험"),
        ("White & Black","음양의 균형","03","대립이 아닌 조화\n동양적 사유의 공간"),
        ("Becoming","농부의 아들","04","9m 스크린 영상 + 청도의 흙\n땅·신체·시간의 순환"),
        ("야외 조각","브론즈 붓질 × 6","05","10m 높이 브론즈 조각\n산세와 건축과의 대화"),
        ("전체 동선","유기적 사유의 흐름","06","안도 다다오의 건축과 함께\n하나의 연결된 서사"),
    ]):
        col, row = i%3, i//3
        cx = sx + col*(cw+gx)
        cy = 1645920 + row*(ch+gy)
        add_rect(slide, cx, cy, cw, ch, fill=RGBColor(0x22,0x22,0x22), line=RGBColor(0x3A,0x3A,0x3A))
        add_rect(slide, cx, cy, 228600, ch, fill=RGBColor(0x2A,0x2A,0x2A))
        add_text(slide, cx+45720, cy, 182880, ch, num, 14, C_GOLD,
                 bold=True, align=PP_ALIGN.CENTER, anchor="ctr")
        add_text(slide, cx+274320, cy+91440, cw-320040, 320040, title, 14, C_WHITE, bold=True)
        add_text(slide, cx+274320, cy+411480, cw-320040, 274320, sub, 10, C_GOLD, italic=True)
        add_text(slide, cx+274320, cy+685800, cw-320040, ch-730000, desc, 11, C_GRAY_L)

# ──────────────────────────────────────────────
# 슬라이드 5 — 기다림
# ──────────────────────────────────────────────
def s5(slide, imgs):
    set_bg_image(slide, imgs["img_4"])
    add_rect(slide, 0, 0, 9144000, 457200, fill=RGBColor(0x1A,0x1A,0x1A), fa=90000)
    add_text(slide, 457200, 0, 8229600, 457200, "T H E M E",
             11, C_GRAY_M, spc=600, anchor="ctr")
    add_rect(slide, 0, 457200, 4500000, 4686300, fill=RGBColor(0x0A,0x0A,0x0A), fa=70000)
    add_text(slide, 457200, 685800, 4200000, 914400,
             "기다림이란 무엇인가", 40, C_WHITE, bold=True)
    add_rect(slide, 548640, 1828800, 54864, 2194560, fill=C_GOLD)
    add_text(slide, 731520, 1828800, 3500000, 2194560,
             '"기다림은 수동적인 것이 아닙니다.\n\n'
             '내가 예술을 진정으로 아는지, 내가 해온 작업은 무엇인지, '
             '더 나아가 작가란 무엇인지 — 그런 마음에서 우러나오는 자기 성찰의 과정을 '
             '나타내는 물리적 시간입니다."',
             14, C_GRAY_L, italic=True)
    add_text(slide, 731520, 4023360, 3500000, 365760, "— 이배",
             14, C_GOLD, italic=True, bold=True)
    tx = 4800000
    add_rect(slide, tx, 1280160, 4100000, 365760, fill=RGBColor(0x1E,0x1E,0x1E), fa=85000)
    add_text(slide, tx, 1280160, 4100000, 365760, "숯이 되기까지의 기다림",
             12, C_GOLD, align=PP_ALIGN.CENTER, anchor="ctr")
    for i, (icon, label, desc) in enumerate([
        ("🌳","나무","살아있는 나무가 베어지기까지의 시간"),
        ("🔥","불","가마 속에서 형태를 잃어가는 연소의 시간"),
        ("⬛","숯","식으며 새로운 물질로 재탄생하는 시간"),
    ]):
        sy = 1645920 + i*1041440
        add_rect(slide, tx, sy, 4100000, 950000, fill=RGBColor(0x18,0x18,0x18), fa=88000,
                 line=RGBColor(0x33,0x33,0x33))
        add_rect(slide, tx+91440, sy+182880, 548640, 548640, fill=RGBColor(0x2A,0x2A,0x2A))
        add_text(slide, tx+91440, sy+182880, 548640, 548640, icon, 20, C_WHITE,
                 align=PP_ALIGN.CENTER, anchor="ctr")
        add_text(slide, tx+731520, sy+182880, 3200000, 320040, label, 18, C_WHITE, bold=True)
        add_text(slide, tx+731520, sy+502920, 3200000, 365760, desc, 12, C_GRAY_L)

# ──────────────────────────────────────────────
# 슬라이드 6 — 엔딩
# ──────────────────────────────────────────────
def s6(slide, imgs):
    set_bg_color(slide, RGBColor(0x1A,0x1A,0x1A))
    add_rect(slide, 0, 0, 9144000, 54864, fill=C_GOLD)
    add_rect(slide, 0, 5088636, 9144000, 54864, fill=C_GOLD)
    add_img(slide, imgs["img_6"], 5200000, 600000, 3700000, 3500000)
    add_img(slide, imgs["img_3"], 5200000, 4100000, 1800000, 900000)
    add_text(slide, 457200, 457200, 4600000, 457200,
             "E N   A T T E N D A N T", 11, C_GRAY_M, spc=600)
    add_text(slide, 457200, 914400, 4600000, 1097280, "기다리며", 80, C_WHITE, bold=True)
    add_lines(slide, 914400, 2103120, 4200000, 822960,
              ["뮤지엄 산 최초 국내 작가 기획전", "이배의 30년을 온전히 담은 전시"],
              16, C_GRAY_L)
    add_rect(slide, 1371600, 3063240, 3500000, 54864, fill=C_GOLD)
    add_rect(slide, 1371600, 3200400, 3400000, 1280160,
             fill=RGBColor(0x22,0x22,0x22), line=RGBColor(0x3A,0x3A,0x3A))
    for i, (icon, txt) in enumerate([
        ("📍 ","강원도 원주  뮤지엄 산 (Museum SAN)"),
        ("📅 ","2026년 4월 7일 — 12월 6일"),
        ("🎨 ","회화 · 조각 · 설치 · 영상  총 39점"),
    ]):
        add_text(slide, 1554480, 3350000+i*355000, 3200000, 320040,
                 f"{icon}{txt}", 13, C_GRAY_L)

# ──────────────────────────────────────────────
# 메인
# ──────────────────────────────────────────────
def build(output="leebae_docent_final.pptx"):
    print("=" * 55)
    print("  이배 《En attendant: 기다리며》 도슨트 PPT 빌드")
    print("=" * 55)
    print("\n이미지 확인 중...")
    imgs = load_images()

    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for label, fn in [
        ("슬라이드 1 — 표지",      s1),
        ("슬라이드 2 — 작가 소개", s2),
        ("슬라이드 3 — 숯의 언어", s3),
        ("슬라이드 4 — 전시 공간", s4),
        ("슬라이드 5 — 기다림",    s5),
        ("슬라이드 6 — 엔딩",      s6),
    ]:
        print(f"  빌드: {label} ...", end="", flush=True)
        fn(prs.slides.add_slide(blank), imgs)
        print(" OK")

    print(f"\n저장: {output}")
    prs.save(output)
    sz = Path(output).stat().st_size
    print(f"완료! {sz:,} bytes  ->  {Path(output).resolve()}")
    print("=" * 55)

if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "leebae_docent_final.pptx")
